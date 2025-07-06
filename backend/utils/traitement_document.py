from deep_translator import GoogleTranslator
import markdown
import PyPDF2
from pdf2image import convert_from_path
import pdfplumber
import langid
import hashlib
import difflib
import re
import pandas as pd
from bs4 import BeautifulSoup
from PIL import Image
import docx2txt
import pytesseract
ALLOWED_EXTENSIONS = {
    "pdf", "txt", "doc", "docx", "rtf", 
    "xls", "xlsx", "csv", "ppt", "pptx", 
    "png", "jpg", "jpeg", "tiff", 
    "xml", "json", "html", "md"
}

def translate_text(text, target_lang):
    try:
        return GoogleTranslator(source='auto', target=target_lang).translate(text)
    except Exception:
        return text
    
def find_best_fuzzy_match(haystack, needle, min_ratio=0.6):
    needle_len = len(needle)
    best_start = -1
    best_ratio = 0
    best_substring = None

    for start in range(len(haystack) - needle_len + 1):
        substring = haystack[start : start + needle_len]
        ratio = difflib.SequenceMatcher(None, needle, substring).ratio()
        if ratio > best_ratio and ratio >= min_ratio:
            best_ratio = ratio
            best_start = start
            best_substring = substring

    if best_start == -1:
        return None, None, None
    else:
        return best_substring, best_start, best_start + len(best_substring)
def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
def compute_file_hash(filepath):
    """Compute SHA256 hash of a file to prevent duplicates"""
    hash_sha256 = hashlib.sha256()
    with open(filepath, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            hash_sha256.update(chunk)
    return hash_sha256.hexdigest()

def clean_text(text):
    """Advanced text cleaning"""
    # Remove non-printable characters
    text = re.sub(r'[\x00-\x1F\x7F-\x9F]', ' ', text)
    
    # Remove common headers/footers
    patterns = [
        r'\b(page|confidential|draft)\b.*\n',
        r'\d{1,2}/\d{1,2}/\d{2,4}',
        r'©.*\n',
        r'http[s]?://\S+',
        r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    ]
    
    for pattern in patterns:
        text = re.sub(pattern, ' ', text, flags=re.IGNORECASE)
    
    # Normalize whitespace
    text = re.sub(r'\s+', ' ', text)
    return text.strip()
def create_thumbnail(filepath, save_path, size=(200, 200)):
    ext = filepath.rsplit('.', 1)[1].lower()
    try:
        if ext == 'pdf':
            # Extraire la première page du PDF comme image
            with pdfplumber.open(filepath) as pdf:
                page = pdf.pages[0]
                pil_image = page.to_image(resolution=150).original
                pil_image.thumbnail(size)
                pil_image.save(save_path)
        elif ext in ['jpg', 'jpeg', 'png', 'gif']:
            img = Image.open(filepath)
            img.thumbnail(size)
            img.save(save_path)
        elif ext in [".docx"]:
            # Pour DOCX, générer une thumbnail textuelle
            text = docx2txt.process(filepath)
            if text:
                img = Image.new("RGB", (400, 300), color=(240, 240, 240))
                from PIL import ImageDraw, ImageFont
                draw = ImageDraw.Draw(img)
                draw.text((10, 10), text[:300] + "...", fill=(0, 0, 0))
                img.save(save_path, "PNG")
        else:
            # Pas de thumbnail pour autres fichiers
            pass
    except Exception as e:
        print(f"Erreur création thumbnail pour {filepath} : {e}")

def extract_text_from_pdf(filepath):
    text = ""
    try:
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    cleaned_text = page_text.strip()
                    text += cleaned_text + "\n"
    except Exception as e:
        print(f"Erreur extraction PDF {filepath} : {e}")
    
    if not text.strip():
        try:
            with open(filepath, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text.strip() + "\n"
        except Exception as e2:
            print(f"PyPDF2 aussi a échoué sur {filepath} : {e2}")
    
    return text.strip()

def extract_text_with_ocr(filepath):
    """OCR extraction with automatic language detection"""
    try:

        
        lang_map = {
            'en': 'eng',
            'fr': 'fra',
            'es': 'spa',
            'de': 'deu',
            'it': 'ita'
        }
        
        text = ""
        if filepath.lower().endswith('.pdf'):
            images = convert_from_path(filepath)
        else:
            images = [Image.open(filepath)]
        
        for img in images:
            # Language detection on sample
            sample_text = pytesseract.image_to_string(img, lang='eng+fra')
            lang, _ = langid.classify(sample_text[:500] if sample_text else 'en')
            lang_code = lang_map.get(lang, 'eng')
            
            # OCR with detected language
            text += pytesseract.image_to_string(img, lang=lang_code) + "\n\n"
        
        return clean_text(text)
    except Exception as e:
        print(f"Erreur OCR: {e}")
        return ""

def extract_office_document(filepath):
    """Extraction robuste pour documents Word"""
    import subprocess
    from docx import Document

    ext = os.path.splitext(filepath)[1].lower()
    text = ""

    try:
        if ext == ".docx":
            # ✅ Extraction avec python-docx
            doc = Document(filepath)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            print(f"[DOCX] Texte extrait depuis {filepath} : {len(text)} caractères")

        elif ext == ".doc":
            # ✅ Extraction avec antiword
            try:
                result = subprocess.run(
                    ["antiword", filepath],
                    capture_output=True,
                    text=True,
                    timeout=30
                )
                text = result.stdout
                print(f"[DOC] Texte extrait via antiword ({len(text)} caractères)")
            except Exception as e:
                print(f"[DOC] Erreur antiword : {e}")
                text = ""

        elif ext == ".rtf":
            # Optionnel : support RTF
            try:
                import striprtf
                with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                    rtf_text = f.read()
                    text = striprtf.rtf_to_text(rtf_text)
            except Exception as e:
                print(f"[RTF] Erreur striprtf : {e}")
                text = ""

        if not text.strip():
            print(f"[⚠️] Aucun texte extrait → fallback OCR")
            return extract_text_with_ocr(filepath)

        return clean_text(text)

    except Exception as e:
        print(f"[Office] ⚠️ Erreur extraction Office : {e}")
        return extract_text_with_ocr(filepath)
def extract_spreadsheet(filepath):
    """Extraction for spreadsheets"""
    try:
        text = ""
        ext = os.path.splitext(filepath)[1].lower()
        
        if ext in ['.xls', '.xlsx']:
            xl = pd.ExcelFile(filepath)
            for sheet_name in xl.sheet_names:
                df = xl.parse(sheet_name)
                text += f"\n\n--- Feuille: {sheet_name} ---\n"
                text += df.to_string(index=False, max_rows=20)
        
        elif ext == '.csv':
            df = pd.read_csv(filepath, nrows=100)  # Limit to 100 rows
            text = df.to_string(index=False)
        
        return clean_text(text)
    except Exception as e:
        print(f"Erreur extraction tableur: {e}")
        return ""

# def extract_presentation(filepath):
#     """Extraction for PowerPoint presentations"""
#     try:
#         from pptx import Presentation
#         text = ""
#         prs = Presentation(filepath)
        
#         for i, slide in enumerate(prs.slides):
#             text += f"\n\n--- Slide {i+1} ---\n"
#             for shape in slide.shapes:
#                 if hasattr(shape, "text"):
#                     text += shape.text + "\n"
        
#         return clean_text(text)
#     except Exception as e:
#         print(f"Erreur extraction présentation: {e}")
#         return extract_text_with_ocr(filepath)

def extract_structured_content(filepath, format_type):
    """Extraction for structured formats"""
    content = ""
    try:
        if format_type == "xml":
            import xml.etree.ElementTree as ET
            tree = ET.parse(filepath)
            root = tree.getroot()
            
            if root.tag.endswith('}article'):  # JATS
                for elem in root.findall('.//abstract'):
                    content += elem.text + "\n"
                for elem in root.findall('.//body'):
                    content += ET.tostring(elem, encoding='unicode', method='text')
            elif root.tag == "TEI":  # TEI XML
                for elem in root.findall('.//text'):
                    content += ET.tostring(elem, encoding='unicode', method='text')
            else:  # Generic XML
                content = ET.tostring(root, encoding='unicode', method='text')
        
        elif format_type == "json":
            import json
            with open(filepath, 'r') as f:
                data = json.load(f)
                
                def extract_values(obj):
                    if isinstance(obj, dict):
                        for v in obj.values():
                            yield from extract_values(v)
                    elif isinstance(obj, list):
                        for item in obj:
                            yield from extract_values(item)
                    else:
                        yield str(obj)
                
                content = " ".join(extract_values(data))
        
        elif format_type == "html":
            with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
                soup = BeautifulSoup(f, 'html.parser')
                content = soup.get_text(separator='\n', strip=True)
        
        elif format_type == "md":
            with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
                md_text = f.read()
                html = markdown.markdown(md_text)
                soup = BeautifulSoup(html, 'html.parser')
                content = soup.get_text(separator='\n', strip=True)
    
    except Exception as e:
        print(f"Erreur extraction structurée: {e}")
    
    return content
from langchain.text_splitter import RecursiveCharacterTextSplitter
import os

def extract_file_content(filepath):
    """Extraction générique du contenu d'un fichier, avec fallback OCR si besoin"""
    ext = os.path.splitext(filepath)[1].lower().replace('.', '')
    content = ""

    try:
        if ext == 'pdf':
            content = extract_text_from_pdf(filepath)
            if not content.strip() or len(content) < 100:
                print("[INFO] Faible contenu PDF – passage en OCR")
                content = extract_text_with_ocr(filepath)

        elif ext in ['png', 'jpg', 'jpeg', 'tiff']:
            content = extract_text_with_ocr(filepath)

        elif ext in ['doc', 'docx', 'rtf']:
            content = extract_office_document(filepath)

        elif ext in ['xls', 'xlsx', 'csv']:
            content = extract_spreadsheet(filepath)

        elif ext in ['xml', 'json', 'html', 'md']:
            content = extract_structured_content(filepath, ext)

        else:  # txt ou inconnu
            try:
                with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                    content = f.read()
            except Exception as e:
                print(f"[ERREUR] Lecture fichier brut : {e}")
                content = extract_text_with_ocr(filepath)

    except Exception as e:
        print(f"[ERREUR] Extraction échouée ({ext}) : {e}")
        content = extract_text_with_ocr(filepath)

    if not content.strip():
        print("[⚠️] Aucun contenu extrait après fallback.")
        return []

    # ✅ Nettoyage et découpage
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200,
        length_function=len,
        add_start_index=True,
    )
    chunks = text_splitter.split_text(clean_text(content))
    print(f"[✅] {len(chunks)} chunks générés à partir du fichier : {os.path.basename(filepath)}")

    return chunks


def calculate_file_hash(filepath):
    sha256_hash = hashlib.sha256()
    with open(filepath, "rb") as f:
        for byte_block in iter(lambda: f.read(4096), b""):
            sha256_hash.update(byte_block)
    return sha256_hash.hexdigest()